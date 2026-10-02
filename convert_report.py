import os
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_styled_paragraph(doc, text="", style=None, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4, line_spacing=1.2):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    pPr = p.paragraph_format
    pPr.space_before = Pt(space_before)
    pPr.space_after = Pt(space_after)
    pPr.line_spacing = line_spacing
    return p

def format_inlines(paragraph, text, base_font_size=13, is_header=False, font_name="Times New Roman"):
    # Tokenize bold, italic, code
    # Regex: `code`, **bold**, *italic*
    tokens = re.split(r'(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_font_size - 1.5)
            run.font.color.rgb = RGBColor(199, 37, 78)
        elif token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.name = font_name
            run.font.size = Pt(base_font_size)
            run.font.bold = True
            if is_header:
                run.font.color.rgb = RGBColor(24, 43, 73)
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = font_name
            run.font.size = Pt(base_font_size)
            run.font.italic = True
        else:
            run = paragraph.add_run(token)
            run.font.name = font_name
            run.font.size = Pt(base_font_size)
            if is_header:
                run.font.color.rgb = RGBColor(24, 43, 73)

def convert_md_to_docx(md_path, docx_path):
    print(f"Reading {md_path}...")
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # Page setup: A4, Vietnamese standard margins
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # 210mm
    section.page_height = Inches(11.69) # 297mm
    section.top_margin = Inches(0.79)   # 20mm
    section.bottom_margin = Inches(0.79)# 20mm
    section.left_margin = Inches(1.18)  # 30mm
    section.right_margin = Inches(0.79) # 20mm

    # Normal style default font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = RGBColor(33, 37, 41)

    in_code_block = False
    code_block_lines = []
    code_block_lang = ""

    in_table = False
    table_lines = []

    align_center = False
    align_right = False

    i = 0
    total_lines = len(lines)

    while i < total_lines:
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()

        # Handle page breaks
        if '<div style="page-break-after: always;"></div>' in stripped:
            doc.add_page_break()
            i += 1
            continue

        # Handle alignment divs
        if '<div align="center">' in stripped:
            align_center = True
            i += 1
            continue
        elif '<div align="right">' in stripped:
            align_right = True
            i += 1
            continue
        elif '</div>' in stripped:
            align_center = False
            align_right = False
            i += 1
            continue

        # Handle horizontal divider
        if stripped in ['---', '***', '___']:
            if not in_code_block:
                # Add a subtle page break or decorative line if needed, or skip
                i += 1
                continue

        # Handle Code blocks
        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_block_lang = stripped[3:].strip()
                code_block_lines = []
            else:
                in_code_block = False
                # Write code block as a single styled table cell with light gray background
                if code_block_lines:
                    tbl = doc.add_table(rows=1, cols=1)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.autofit = False
                    cell = tbl.cell(0, 0)
                    cell.width = Inches(6.3)
                    set_cell_background(cell, "F8F9FA")
                    
                    # Add code content
                    code_text = '\n'.join(code_block_lines)
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_before = Pt(2)
                    cp.paragraph_format.space_after = Pt(2)
                    cp.paragraph_format.line_spacing = 1.0
                    crun = cp.add_run(code_text)
                    crun.font.name = "Consolas"
                    crun.font.size = Pt(9.5)
                    crun.font.color.rgb = RGBColor(40, 44, 52)
                    
                    # Empty space after table
                    sp = doc.add_paragraph()
                    sp.paragraph_format.space_before = Pt(0)
                    sp.paragraph_format.space_after = Pt(4)
                code_block_lines = []
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Handle Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            # Check if next line is also a table row
            if i + 1 < total_lines and lines[i+1].strip().startswith('|') and lines[i+1].strip().endswith('|'):
                i += 1
                continue
            else:
                # Process table
                rows = []
                for tline in table_lines:
                    # check if separator row like |---|---|
                    cells = [c.strip() for c in tline.strip('|').split('|')]
                    if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                        continue
                    rows.append(cells)

                if rows:
                    col_count = max(len(r) for r in rows)
                    tbl = doc.add_table(rows=len(rows), cols=col_count)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(tbl, "B0C4DE")

                    for r_idx, row_data in enumerate(rows):
                        for c_idx in range(col_count):
                            val = row_data[c_idx] if c_idx < len(row_data) else ""
                            cell = tbl.cell(r_idx, c_idx)
                            cp = cell.paragraphs[0]
                            cp.paragraph_format.space_before = Pt(1)
                            cp.paragraph_format.space_after = Pt(2)
                            cp.paragraph_format.line_spacing = 1.02

                            if r_idx == 0:
                                set_cell_background(cell, "EBF3FB")
                                format_inlines(cp, val, base_font_size=11, is_header=True)
                                cp.runs[0].font.bold = True if cp.runs else None
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_background(cell, "FFFFFF")
                                else:
                                    set_cell_background(cell, "F7FAFD")
                                format_inlines(cp, val, base_font_size=10)

                    # Space after table
                    sp = doc.add_paragraph()
                    sp.paragraph_format.space_before = Pt(0)
                    sp.paragraph_format.space_after = Pt(6)

                table_lines = []
                i += 1
                continue

        # Skip empty lines
        if not stripped or stripped == '<br/>':
            i += 1
            continue

        # Handle Headings
        current_align = WD_ALIGN_PARAGRAPH.CENTER if align_center else (WD_ALIGN_PARAGRAPH.RIGHT if align_right else WD_ALIGN_PARAGRAPH.LEFT)

        if stripped.startswith('# '):
            h_text = stripped[2:].strip()
            p = add_styled_paragraph(doc, align=current_align, space_before=10, space_after=5, line_spacing=1.1)
            format_inlines(p, h_text, base_font_size=16, is_header=True)
            for run in p.runs:
                run.font.bold = True
        elif stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = add_styled_paragraph(doc, align=current_align, space_before=7, space_after=4, line_spacing=1.1)
            format_inlines(p, h_text, base_font_size=14, is_header=True)
            for run in p.runs:
                run.font.bold = True
        elif stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = add_styled_paragraph(doc, align=current_align, space_before=5, space_after=3, line_spacing=1.1)
            format_inlines(p, h_text, base_font_size=13, is_header=True)
            for run in p.runs:
                run.font.bold = True
        elif stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = add_styled_paragraph(doc, align=current_align, space_before=4, space_after=2, line_spacing=1.1)
            format_inlines(p, h_text, base_font_size=13, is_header=True)
            for run in p.runs:
                run.font.italic = True
        elif stripped.startswith('- ') or stripped.startswith('* '):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.12
            format_inlines(p, bullet_text, base_font_size=13)
        elif re.match(r'^\d+\.\s+', stripped):
            num_match = re.match(r'^(\d+\.)\s+(.*)', stripped)
            prefix = num_match.group(1)
            num_text = num_match.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.12
            prun = p.add_run(prefix + " ")
            prun.font.name = "Times New Roman"
            prun.font.size = Pt(13)
            prun.font.bold = True
            format_inlines(p, num_text, base_font_size=13)
        elif stripped.startswith('> '):
            quote_text = stripped[2:].strip()
            p = add_styled_paragraph(doc, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=4, space_after=4, line_spacing=1.2)
            p.paragraph_format.left_indent = Inches(0.4)
            format_inlines(p, quote_text, base_font_size=12.5)
            for run in p.runs:
                run.font.italic = True
                run.font.color.rgb = RGBColor(108, 117, 125)
        else:
            # Regular paragraph
            p = add_styled_paragraph(doc, align=current_align, space_before=1, space_after=4, line_spacing=1.18)
            format_inlines(p, stripped, base_font_size=13)

        i += 1

    print(f"Saving to {docx_path}...")
    doc.save(docx_path)
    print("Done!")

if __name__ == '__main__':
    md_file = sys.argv[1]
    docx_file = sys.argv[2]
    convert_md_to_docx(md_file, docx_file)
